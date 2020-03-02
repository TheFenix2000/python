from docx import Document

document = Document('header.docx')
document.add_page_break()

def Make_Table(rows_number):
    table = document.add_table(cols=1, rows=rows_number, style='Table Grid')
    for i in range(0,rows_number):
        cell = table.cell(i,0)
        cell.text = "Hi"
        i += 1
    document.save("Nessus-result.docx")
Make_Table(7)
